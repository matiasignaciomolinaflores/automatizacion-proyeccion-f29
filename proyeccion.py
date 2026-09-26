from __future__ import annotations

import base64
import logging
import re
import sys
import time
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from email.message import EmailMessage
from pathlib import Path
from typing import Dict, Iterable, Optional

import requests
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait


# ============================================================
# CONFIGURACIÓN
# ============================================================
LOGIN_URL = (
    "https://zeusr.sii.cl/AUT2000/InicioAutenticacion/IngresoRutClave.html?"
    "https://www4.sii.cl/consdcvinternetui/"
)
F29_URL = "https://www4.sii.cl/rfiInternet/consulta/index.html#rfiSelFormularioPeriodo"

ARCHIVO_EXCEL_PREDETERMINADO = "plantilla definitiva.xlsx"
NOMBRE_HOJA = "Hoja1"
FILA_INICIAL = 5

# Plantilla real
COL_NOMBRE_CLIENTE = "C"
COL_CORREO = "D"
COL_RAZON_SOCIAL = "E"
COL_RUT = "F"
COL_CLAVE = "G"
COL_IVA_DEBITO = "H"
COL_IVA_CREDITO = "I"
COL_REMANENTE = "J"
COL_IVA_CREDITO_TOTAL = "O"
COL_RESULTADO_IVA = "P"
COL_IVA_POR_PAGAR = "Q"
COL_REMANENTE_SIGUIENTE = "R"
COL_RBH = "S"
COL_IU = "T"
COL_BASE_IMPONIBLE = "U"
COL_TASA_PPM = "V"
COL_PPM = "W"
COL_TOTAL_F29 = "X"

CORREO_REMITENTE = "contabilidadespvm@gmail.com"
ARCHIVO_CREDENCIALES_GOOGLE = "credentials.json"
ARCHIVO_TOKEN_GOOGLE = "token.json"
GMAIL_SCOPES = ["https://www.googleapis.com/auth/gmail.send"]

TIMEOUT = 35
TIMEOUT_DESCARGA = 45

MESES_ES = {
    1: "Enero",
    2: "Febrero",
    3: "Marzo",
    4: "Abril",
    5: "Mayo",
    6: "Junio",
    7: "Julio",
    8: "Agosto",
    9: "Septiembre",
    10: "Octubre",
    11: "Noviembre",
    12: "Diciembre",
}


# ============================================================
# EXCEPCIONES
# ============================================================
class AutomatizacionError(Exception):
    pass


class LoginError(AutomatizacionError):
    pass


class RutNoDisponibleError(AutomatizacionError):
    pass


class TablaRCVError(AutomatizacionError):
    pass


class NoF29Error(AutomatizacionError):
    pass


class F29ValidacionError(AutomatizacionError):
    pass


class F29ExtraccionError(AutomatizacionError):
    pass


# ============================================================
# RESULTADOS
# ============================================================
@dataclass
class ResultadoRCV:
    iva_credito: int
    iva_debito: int
    base_imponible: int


@dataclass
class ResultadoF29:
    remanente_077: int
    rbh_151: int
    iu_048: int
    tasa_ppm_115: Decimal  # Ej.: Decimal("2") significa 2 %
    rut_pdf: str
    periodo_pdf: str


# ============================================================
# LOG
# ============================================================
def configurar_log(carpeta: Path) -> None:
    ruta_log = carpeta / "automatizacion_sii.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[
            logging.FileHandler(ruta_log, encoding="utf-8"),
        ],
    )


def log_info(mensaje: str) -> None:
    print(mensaje)
    logging.info(mensaje)


def log_error(mensaje: str) -> None:
    print(mensaje)
    logging.error(mensaje)


# ============================================================
# UTILIDADES GENERALES
# ============================================================
def normalizar_texto(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    texto = re.sub(r"\s+", " ", texto).strip().upper()
    return texto


def texto_excel(valor: object) -> str:
    """Convierte una celda Excel a texto sin dejar '123.0' cuando era entero."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def normalizar_rut(rut: object) -> str:
    return texto_excel(rut).upper().replace(".", "").replace(" ", "").strip()


def nombre_archivo_seguro(nombre: str) -> str:
    nombre = re.sub(r'[\\/:*?"<>|]', "_", nombre).strip().rstrip(".")
    nombre = re.sub(r"\s+", " ", nombre)
    return nombre or "cliente"


def monto_entero_chile(valor: object) -> int:
    """'30.139.134', '$ 409.370', '0', '(10.000)' -> int."""
    if valor is None:
        return 0

    texto = str(valor).strip()
    if not texto or texto in {"-", "—"}:
        return 0

    negativo = texto.startswith("-") or (texto.startswith("(") and texto.endswith(")"))
    solo_digitos = re.sub(r"\D", "", texto)
    if not solo_digitos:
        return 0

    numero = int(solo_digitos)
    return -numero if negativo else numero


def decimal_chile(valor: object) -> Decimal:
    """'2', '2,5', '0,25' -> Decimal. NO divide por 100 aquí."""
    if valor is None:
        return Decimal("0")

    texto = str(valor).strip().replace("%", "").replace(" ", "")
    if not texto:
        return Decimal("0")

    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")

    try:
        return Decimal(texto)
    except InvalidOperation as exc:
        raise F29ExtraccionError(f"No se pudo convertir el valor decimal {valor!r}.") from exc


def periodo_anterior(anio: int, mes: int) -> tuple[int, int]:
    if mes == 1:
        return anio - 1, 12
    return anio, mes - 1


def periodo_yyyymm(anio: int, mes: int) -> str:
    return f"{anio:04d}{mes:02d}"


def imprimir_cliente(fila: int, rut: str, nombre: str) -> None:
    log_info("\n" + "=" * 60)
    log_info(f"CLIENTE FILA {fila}")
    log_info(f"NOMBRE: {nombre}")
    log_info(f"RUT: {rut}")
    log_info("=" * 60)


# ============================================================
# EXCEL
# ============================================================
def abrir_excel(ruta: Path):
    if not ruta.exists():
        raise FileNotFoundError(f"No existe el Excel: {ruta}")

    wb = load_workbook(ruta, data_only=False)
    if NOMBRE_HOJA not in wb.sheetnames:
        raise AutomatizacionError(
            f"No existe la hoja {NOMBRE_HOJA!r}. Hojas disponibles: {wb.sheetnames}"
        )
    return wb, wb[NOMBRE_HOJA]


def guardar_excel(wb, ruta: Path) -> None:
    try:
        wb.save(ruta)
    except PermissionError as exc:
        raise AutomatizacionError(
            "No se pudo guardar el Excel. Ciérralo en Microsoft Excel y vuelve a ejecutar."
        ) from exc


def escribir_resultados(ws, fila: int, rcv: ResultadoRCV, f29: ResultadoF29) -> None:
    ws[f"{COL_IVA_DEBITO}{fila}"] = rcv.iva_debito
    ws[f"{COL_IVA_CREDITO}{fila}"] = rcv.iva_credito
    ws[f"{COL_REMANENTE}{fila}"] = f29.remanente_077
    ws[f"{COL_RBH}{fila}"] = f29.rbh_151
    ws[f"{COL_IU}{fila}"] = f29.iu_048
    ws[f"{COL_BASE_IMPONIBLE}{fila}"] = rcv.base_imponible

    # El código 115 del F29 viene como "2" para indicar 2 %.
    # Excel almacena 2 % como 0.02. Así W = U * V funciona correctamente.
    celda_tasa = ws[f"{COL_TASA_PPM}{fila}"]
    celda_tasa.value = float(f29.tasa_ppm_115 / Decimal("100"))
    celda_tasa.number_format = "0.##%"


# ============================================================
# CHROME / SELENIUM
# ============================================================
def crear_driver(carpeta_descargas: Path) -> webdriver.Chrome:
    options = webdriver.ChromeOptions()
    prefs = {
        "download.default_directory": str(carpeta_descargas.resolve()),
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        # Dejamos el visor PDF de Chrome habilitado. Luego descargamos el mismo
        # recurso desde Python usando las cookies de la sesión de Selenium.
        "plugins.always_open_pdf_externally": False,
        "safebrowsing.enabled": True,
    }
    options.add_experimental_option("prefs", prefs)
    options.add_argument("--start-maximized")

    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(60)

    # Refuerzo del directorio de descarga para versiones recientes de Chrome.
    try:
        driver.execute_cdp_cmd(
            "Page.setDownloadBehavior",
            {"behavior": "allow", "downloadPath": str(carpeta_descargas.resolve())},
        )
    except Exception:
        pass

    return driver


def iniciar_sesion(driver, rut: str, clave: str) -> None:
    wait = WebDriverWait(driver, TIMEOUT)
    driver.get(LOGIN_URL)

    campo_rut = wait.until(EC.visibility_of_element_located((By.ID, "rutcntr")))
    campo_clave = wait.until(EC.visibility_of_element_located((By.ID, "clave")))

    campo_rut.clear()
    campo_rut.send_keys(rut)
    campo_clave.clear()
    campo_clave.send_keys(clave)

    wait.until(EC.element_to_be_clickable((By.ID, "bt_ingresar"))).click()

    try:
        wait.until(EC.presence_of_element_located((By.XPATH, "//select[@ng-model='rut']")))
    except TimeoutException as exc:
        raise LoginError(
            "No se llegó a la pantalla del Registro de Compras y Ventas. "
            "Revisa RUT/clave o si el SII mostró un mensaje de error."
        ) from exc


# ============================================================
# REGISTRO DE COMPRAS Y VENTAS
# ============================================================
def seleccionar_rut_empresa(driver, rut_cliente: str) -> None:
    wait = WebDriverWait(driver, TIMEOUT)
    elemento = wait.until(
        EC.presence_of_element_located((By.XPATH, "//select[@ng-model='rut']"))
    )
    selector = Select(elemento)
    rut_objetivo = normalizar_rut(rut_cliente)

    for opcion in selector.options:
        valor = opcion.get_attribute("value") or opcion.text
        if normalizar_rut(valor) == rut_objetivo:
            selector.select_by_visible_text(opcion.text)
            seleccionado = normalizar_rut(selector.first_selected_option.text)
            if seleccionado != rut_objetivo:
                raise RutNoDisponibleError(
                    f"RUT seleccionado {seleccionado} distinto de {rut_objetivo}."
                )
            return

    raise RutNoDisponibleError(
        f"El RUT {rut_cliente} no aparece entre las empresas disponibles del usuario."
    )


def seleccionar_periodo_rcv(driver, anio: int, mes: int) -> None:
    wait = WebDriverWait(driver, TIMEOUT)

    selector_mes = Select(
        wait.until(EC.presence_of_element_located((By.ID, "periodoMes")))
    )
    selector_anio = Select(
        wait.until(
            EC.presence_of_element_located(
                (By.XPATH, "//select[@ng-model='periodoAnho']")
            )
        )
    )

    selector_mes.select_by_value(f"{mes:02d}")
    selector_anio.select_by_value(str(anio))

    wait.until(
        EC.element_to_be_clickable((By.XPATH, "//button[normalize-space()='Consultar']"))
    ).click()

    wait.until(
        EC.presence_of_element_located(
            (
                By.XPATH,
                "//*[contains(normalize-space(), 'RESUMEN REGISTRO DE COMPRAS')]",
            )
        )
    )


def encabezados_tabla(tabla) -> list[str]:
    celdas = tabla.find_elements(By.XPATH, ".//thead//th")
    if not celdas:
        celdas = tabla.find_elements(By.XPATH, ".//tr[1]/*[self::th or self::td]")
    return [normalizar_texto(c.text) for c in celdas]


def buscar_tabla_con_columnas(driver, columnas_requeridas: Iterable[str]):
    requeridas = {normalizar_texto(c) for c in columnas_requeridas}
    limite = time.time() + TIMEOUT

    while time.time() < limite:
        tablas = driver.find_elements(By.TAG_NAME, "table")
        for tabla in tablas:
            try:
                if not tabla.is_displayed():
                    continue
                headers = set(encabezados_tabla(tabla))
                if requeridas.issubset(headers):
                    return tabla
            except Exception:
                continue
        time.sleep(0.4)

    raise TablaRCVError(
        f"No encontré una tabla visible con las columnas: {', '.join(columnas_requeridas)}"
    )


def leer_tabla_por_encabezados(tabla) -> list[Dict[str, str]]:
    headers = encabezados_tabla(tabla)
    if not headers:
        raise TablaRCVError("La tabla del SII no tiene encabezados legibles.")

    filas = tabla.find_elements(By.XPATH, ".//tbody/tr")
    if not filas:
        filas = tabla.find_elements(By.XPATH, ".//tr[position()>1]")

    resultado: list[Dict[str, str]] = []
    for fila in filas:
        celdas = fila.find_elements(By.XPATH, "./td")
        if not celdas:
            continue
        valores = [c.text.strip() for c in celdas]
        if len(valores) < len(headers):
            continue
        resultado.append({headers[i]: valores[i] for i in range(len(headers))})

    if not resultado:
        raise TablaRCVError("La tabla resumen del SII no contiene filas de datos.")
    return resultado


def obtener_columna(fila: Dict[str, str], columna: str) -> str:
    buscada = normalizar_texto(columna)
    for encabezado, valor in fila.items():
        if normalizar_texto(encabezado) == buscada:
            return valor
    raise TablaRCVError(f"No se encontró la columna {columna!r} en la tabla.")


def es_nota_credito(tipo_documento: str) -> bool:
    tipo = normalizar_texto(tipo_documento)
    return "NOTA DE CREDITO" in tipo or bool(re.search(r"\(\s*61\s*\)", tipo))


def calcular_compras(driver) -> int:
    log_info("[COMPRAS] Leyendo resumen...")
    tabla = buscar_tabla_con_columnas(driver, ["Tipo Documento", "IVA Recuperable"])
    filas = leer_tabla_por_encabezados(tabla)

    iva_credito = 0
    for fila in filas:
        tipo = obtener_columna(fila, "Tipo Documento")
        iva = monto_entero_chile(obtener_columna(fila, "IVA Recuperable"))
        signo = -1 if es_nota_credito(tipo) else 1
        iva_credito += signo * iva
        log_info(
            f"[COMPRAS] {tipo} | IVA Recuperable={iva} | "
            f"{'RESTA' if signo < 0 else 'SUMA'}"
        )

    log_info(f"[COMPRAS] IVA Crédito Fiscal: ${iva_credito:,}".replace(",", "."))
    return iva_credito


def ir_a_ventas(driver) -> None:
    wait = WebDriverWait(driver, TIMEOUT)

    log_info("[VENTAS] Buscando pestaña VENTA...")

    try:
        enlace_venta = wait.until(
            EC.element_to_be_clickable(
                (
                    By.CSS_SELECTOR,
                    'a[ui-sref="venta"]',
                )
            )
        )

        log_info("[VENTAS] Pestaña VENTA encontrada.")
        log_info("[VENTAS] Haciendo click...")

        try:
            enlace_venta.click()
        except Exception:
            driver.execute_script("arguments[0].click();", enlace_venta)

        log_info("[VENTAS] Click realizado.")
        log_info("[VENTAS] Esperando tabla del Registro de Ventas...")

        # El SII puede mantener la URL #/index aunque ya haya cambiado
        # visualmente a la pestaña VENTA. Por eso validamos el contenido
        # real de Ventas en vez de depender de la URL.
        buscar_tabla_con_columnas(
            driver,
            ["Tipo Documento", "Monto Exento", "Monto Neto", "Monto IVA"],
        )

        log_info("[VENTAS] Registro de Ventas cargado correctamente.")

    except TablaRCVError:
        log_error(f"[VENTAS] URL actual al fallar: {driver.current_url}")
        raise

    except TimeoutException as exc:
        log_error(f"[VENTAS] URL actual al fallar: {driver.current_url}")
        raise TablaRCVError(
            "No fue posible encontrar o hacer click en la pestaña VENTA."
        ) from exc


def calcular_ventas(driver) -> tuple[int, int]:
    log_info("[VENTAS] Leyendo resumen...")
    tabla = buscar_tabla_con_columnas(
        driver,
        ["Tipo Documento", "Monto Exento", "Monto Neto", "Monto IVA"],
    )
    filas = leer_tabla_por_encabezados(tabla)

    iva_debito = 0
    base_imponible = 0

    for fila in filas:
        tipo = obtener_columna(fila, "Tipo Documento")
        monto_exento = monto_entero_chile(obtener_columna(fila, "Monto Exento"))
        monto_neto = monto_entero_chile(obtener_columna(fila, "Monto Neto"))
        monto_iva = monto_entero_chile(obtener_columna(fila, "Monto IVA"))

        signo = -1 if es_nota_credito(tipo) else 1
        iva_debito += signo * monto_iva
        base_imponible += signo * (monto_exento + monto_neto)

        log_info(
            f"[VENTAS] {tipo} | Exento={monto_exento} | Neto={monto_neto} | "
            f"IVA={monto_iva} | {'RESTA' if signo < 0 else 'SUMA'}"
        )

    log_info(f"[VENTAS] IVA Débito Fiscal: ${iva_debito:,}".replace(",", "."))
    log_info(f"[VENTAS] Base Imponible: ${base_imponible:,}".replace(",", "."))
    return iva_debito, base_imponible


def procesar_rcv(driver, rut: str, anio: int, mes: int) -> ResultadoRCV:
    seleccionar_rut_empresa(driver, rut)
    seleccionar_periodo_rcv(driver, anio, mes)

    iva_credito = calcular_compras(driver)
    ir_a_ventas(driver)
    iva_debito, base_imponible = calcular_ventas(driver)

    return ResultadoRCV(
        iva_credito=iva_credito,
        iva_debito=iva_debito,
        base_imponible=base_imponible,
    )


# ============================================================
# F29 - NAVEGACIÓN
# ============================================================
def seleccionar_f29_y_periodo(driver, anio: int, mes: int) -> None:
    wait = WebDriverWait(driver, TIMEOUT)
    driver.get(F29_URL)

    wait.until(lambda d: len(d.find_elements(By.CSS_SELECTOR, "select.gwt-ListBox")) >= 3)
    selects = driver.find_elements(By.CSS_SELECTOR, "select.gwt-ListBox")

    sel_formulario: Optional[Select] = None
    sel_anio: Optional[Select] = None
    sel_mes: Optional[Select] = None

    for elemento in selects:
        selector = Select(elemento)
        textos = {o.text.strip() for o in selector.options}
        if "Formulario 29" in textos:
            sel_formulario = selector
        if str(anio) in textos:
            sel_anio = selector
        if MESES_ES[mes] in textos:
            sel_mes = selector

    if not sel_formulario or not sel_anio or not sel_mes:
        raise AutomatizacionError("No se pudieron identificar los selectores del F29.")

    sel_formulario.select_by_visible_text("Formulario 29")
    sel_anio.select_by_visible_text(str(anio))
    sel_mes.select_by_visible_text(MESES_ES[mes])

    wait.until(
        EC.element_to_be_clickable(
            (By.XPATH, "//button[normalize-space()='Buscar Datos Ingresados']")
        )
    ).click()

    wait.until(
        EC.presence_of_element_located(
            (
                By.XPATH,
                "//*[contains(normalize-space(), 'RESULTADOS DE LA BÚSQUEDA') "
                "or contains(normalize-space(), 'RESULTADOS DE LA BUSQUEDA')]",
            )
        )
    )


def obtener_folio_vigente(driver) -> str:
    limite = time.time() + TIMEOUT

    while time.time() < limite:
        tablas = driver.find_elements(By.TAG_NAME, "table")
        for tabla in tablas:
            try:
                texto = normalizar_texto(tabla.text)
                if "DECLARACIONES VIGENTES" not in texto:
                    continue

                enlaces = tabla.find_elements(By.XPATH, ".//a[normalize-space()]")
                for enlace in enlaces:
                    folio = enlace.text.strip()
                    if re.fullmatch(r"\d+", folio):
                        driver.execute_script("arguments[0].click();", enlace)
                        return folio
            except Exception:
                continue
        time.sleep(0.4)

    raise NoF29Error(
        "No aparece un folio clickeable en DECLARACIONES VIGENTES. "
        "Se considera que no existe F29 para ese período."
    )


def abrir_formulario_compacto(driver) -> str:
    """Abre Formulario Compacto y devuelve la URL del recurso en la nueva ventana."""
    wait = WebDriverWait(driver, TIMEOUT)
    ventanas_antes = set(driver.window_handles)

    boton = wait.until(
        EC.element_to_be_clickable(
            (By.XPATH, "//button[normalize-space()='Formulario Compacto']")
        )
    )
    driver.execute_script("arguments[0].click();", boton)

    # Normalmente abre una nueva ventana/pestaña.
    try:
        wait.until(lambda d: len(set(d.window_handles) - ventanas_antes) >= 1)
        nueva = list(set(driver.window_handles) - ventanas_antes)[0]
        driver.switch_to.window(nueva)
    except TimeoutException:
        pass

    wait.until(lambda d: "formCompacto" in d.current_url or d.current_url.startswith("blob:"))
    return driver.current_url


# ============================================================
# F29 - DESCARGA
# ============================================================
def sesion_requests_desde_selenium(driver) -> requests.Session:
    session = requests.Session()
    for cookie in driver.get_cookies():
        session.cookies.set(
            cookie["name"],
            cookie["value"],
            domain=cookie.get("domain"),
            path=cookie.get("path", "/"),
        )

    try:
        ua = driver.execute_script("return navigator.userAgent")
        if ua:
            session.headers.update({"User-Agent": ua})
    except Exception:
        pass

    return session


def descargar_pdf_desde_url_sesion(driver, url: str, destino: Path) -> bool:
    """Descarga el PDF usando las cookies del navegador. Retorna True si obtuvo un PDF."""
    if not url.startswith("http"):
        return False

    session = sesion_requests_desde_selenium(driver)
    try:
        respuesta = session.get(url, timeout=45, allow_redirects=True)
    except requests.RequestException:
        return False

    if respuesta.status_code != 200:
        return False

    contenido = respuesta.content
    content_type = respuesta.headers.get("Content-Type", "").lower()
    es_pdf = contenido.startswith(b"%PDF") or "application/pdf" in content_type
    if not es_pdf:
        return False

    destino.write_bytes(contenido)
    return True


def esperar_pdf_nuevo(carpeta: Path, existentes: set[Path], timeout: int) -> Optional[Path]:
    limite = time.time() + timeout
    existentes_resueltos = {p.resolve() for p in existentes}

    while time.time() < limite:
        temporales = list(carpeta.glob("*.crdownload"))
        actuales = {p.resolve() for p in carpeta.glob("*.pdf")}
        nuevos = actuales - existentes_resueltos

        if nuevos and not temporales:
            return max(nuevos, key=lambda p: p.stat().st_mtime)
        time.sleep(0.5)

    return None


def descargar_f29(
    driver,
    carpeta: Path,
    nombre_cliente: str,
    rut: str,
    anio: int,
    mes: int,
) -> Path:
    log_info(f"[F29] Buscando período {periodo_yyyymm(anio, mes)}...")
    seleccionar_f29_y_periodo(driver, anio, mes)

    folio = obtener_folio_vigente(driver)
    log_info(f"[F29] Folio vigente: {folio}")

    url_compacto = abrir_formulario_compacto(driver)
    destino = carpeta / f"{nombre_archivo_seguro(nombre_cliente)}.pdf"

    # Si ya existía un PDF de una ejecución anterior, lo reemplazaremos.
    if destino.exists():
        destino.unlink()

    # Método principal: descarga directa del recurso usando la misma sesión autenticada.
    if descargar_pdf_desde_url_sesion(driver, url_compacto, destino):
        log_info(f"[F29] PDF descargado: {destino.name}")
        return destino

    # Respaldo: si Chrome, por la configuración, descargó automáticamente el PDF.
    existentes = {p.resolve() for p in carpeta.glob("*.pdf")}
    descargado = esperar_pdf_nuevo(carpeta, existentes, TIMEOUT_DESCARGA)
    if descargado:
        if descargado.resolve() != destino.resolve():
            descargado.replace(destino)
        log_info(f"[F29] PDF descargado: {destino.name}")
        return destino

    raise AutomatizacionError(
        "No fue posible descargar el Formulario Compacto como PDF. "
        f"URL observada: {url_compacto}"
    )


# ============================================================
# F29 - LECTURA Y VALIDACIÓN
# ============================================================
def extraer_texto_pdf(ruta_pdf: Path) -> str:
    try:
        reader = PdfReader(str(ruta_pdf))
        texto = "\n".join((pagina.extract_text() or "") for pagina in reader.pages).strip()
    except Exception as exc:
        raise F29ValidacionError(f"No se pudo abrir/leer el PDF: {exc}") from exc

    if not texto:
        raise F29ValidacionError("El PDF no contiene texto extraíble.")
    return texto


def extraer_rut_pdf(texto: str) -> str:
    m = re.search(r"RUT\s*\[03\]\s*([0-9.\-Kk]+)", texto, flags=re.IGNORECASE)
    if not m:
        raise F29ValidacionError("No se encontró RUT [03] en el F29.")
    return m.group(1).strip()


def extraer_periodo_pdf(texto: str) -> str:
    m = re.search(r"PERIODO\s*\[15\]\s*(\d{6})", texto, flags=re.IGNORECASE)
    if not m:
        raise F29ValidacionError("No se encontró PERIODO [15] en el F29.")
    return m.group(1)


def extraer_codigo(texto: str, codigo: str, decimal: bool = False):
    """
    Si el código NO aparece en un F29 válido -> 0.
    Si aparece pero no se puede obtener su valor -> ERROR.
    """
    patron = re.compile(rf"(?m)^\s*{re.escape(codigo)}\b.*$", flags=re.IGNORECASE)
    coincidencia = patron.search(texto)

    if not coincidencia:
        return Decimal("0") if decimal else 0

    linea = coincidencia.group(0).strip()
    m_valor = re.search(r"([0-9][0-9.,]*)\s*$", linea)
    if not m_valor:
        raise F29ExtraccionError(
            f"El código {codigo} aparece, pero no se pudo extraer su valor. Línea: {linea!r}"
        )

    valor = m_valor.group(1)
    return decimal_chile(valor) if decimal else monto_entero_chile(valor)


def leer_y_validar_f29(
    ruta_pdf: Path,
    rut_esperado: str,
    periodo_esperado: str,
) -> ResultadoF29:
    texto = extraer_texto_pdf(ruta_pdf)

    if "FORMULARIO 29" not in normalizar_texto(texto):
        raise F29ValidacionError("El archivo descargado no corresponde a un FORMULARIO 29.")

    rut_pdf = extraer_rut_pdf(texto)
    if normalizar_rut(rut_pdf) != normalizar_rut(rut_esperado):
        raise F29ValidacionError(
            f"RUT del PDF ({rut_pdf}) distinto del cliente ({rut_esperado})."
        )

    periodo_pdf = extraer_periodo_pdf(texto)
    if periodo_pdf != periodo_esperado:
        raise F29ValidacionError(
            f"Período del PDF ({periodo_pdf}) distinto del esperado ({periodo_esperado})."
        )

    resultado = ResultadoF29(
        remanente_077=extraer_codigo(texto, "077"),
        rbh_151=extraer_codigo(texto, "151"),
        iu_048=extraer_codigo(texto, "048"),
        tasa_ppm_115=extraer_codigo(texto, "115", decimal=True),
        rut_pdf=rut_pdf,
        periodo_pdf=periodo_pdf,
    )

    log_info(f"[F29] Código 077: ${resultado.remanente_077:,}".replace(",", "."))
    log_info(f"[F29] Código 151: ${resultado.rbh_151:,}".replace(",", "."))
    log_info(f"[F29] Código 048: ${resultado.iu_048:,}".replace(",", "."))
    log_info(f"[F29] Código 115 - Tasa PPM: {resultado.tasa_ppm_115}%")
    return resultado



# ============================================================
# UTM - ACTUALIZACIÓN FINAL DE D2 Y G2
# ============================================================
def obtener_utm(anio_consulta: int, mes_consulta: int) -> int:
    url = (
        f"https://www.sii.cl/valores_y_fechas/"
        f"utm/utm{anio_consulta}.htm"
    )

    log_info(f"[UTM] Consultando SII: {url}")

    tablas = pd.read_html(url)

    tabla_utm = None
    for tabla in tablas:
        columnas = " ".join(str(c) for c in tabla.columns)
        if "UTM" in columnas.upper():
            tabla_utm = tabla
            break

    if tabla_utm is None:
        raise AutomatizacionError(
            f"No se encontró la tabla UTM del año {anio_consulta}."
        )

    nombre_mes = MESES_ES[mes_consulta]
    primera_columna = tabla_utm.columns[0]

    fila = tabla_utm[
        tabla_utm[primera_columna]
        .astype(str)
        .str.strip()
        .str.lower()
        == nombre_mes.lower()
    ]

    if fila.empty:
        raise AutomatizacionError(
            f"No se encontró {nombre_mes} {anio_consulta} en el SII."
        )

    valor = fila.iloc[0, 1]

    if pd.isna(valor):
        raise AutomatizacionError(
            f"El SII todavía no ha publicado la UTM de "
            f"{nombre_mes} {anio_consulta}."
        )

    valor_texto = str(valor).strip()
    valor_texto = (
        valor_texto
        .replace("$", "")
        .replace(".", "")
        .replace(",", "")
        .replace(" ", "")
    )

    try:
        return int(valor_texto)
    except ValueError as exc:
        raise AutomatizacionError(
            f"No se pudo convertir la UTM de {nombre_mes} {anio_consulta}: "
            f"{valor!r}"
        ) from exc


def periodo_siguiente(anio: int, mes: int) -> tuple[int, int]:
    if mes == 12:
        return anio + 1, 1
    return anio, mes + 1


def actualizar_utm_excel(
    wb,
    ws,
    ruta_excel: Path,
    anio_actual: int,
    mes_actual: int,
) -> None:
    log_info("\n[UTM] Iniciando actualización final de UTM...")

    anio_siguiente, mes_siguiente = periodo_siguiente(anio_actual, mes_actual)

    utm_actual = obtener_utm(anio_actual, mes_actual)
    utm_siguiente = obtener_utm(anio_siguiente, mes_siguiente)

    log_info(
        f"[UTM] {MESES_ES[mes_actual]} {anio_actual}: "
        f"${utm_actual:,}".replace(",", ".")
    )
    log_info(
        f"[UTM] {MESES_ES[mes_siguiente]} {anio_siguiente}: "
        f"${utm_siguiente:,}".replace(",", ".")
    )

    ws["D2"] = utm_actual
    ws["G2"] = utm_siguiente

    ws["D2"].number_format = "#,##0"
    ws["G2"].number_format = "#,##0"

    guardar_excel(wb, ruta_excel)

    log_info(
        f"[UTM] D2 actualizado con UTM {MESES_ES[mes_actual]} "
        f"{anio_actual}: {utm_actual:,}".replace(",", ".")
    )
    log_info(
        f"[UTM] G2 actualizado con UTM {MESES_ES[mes_siguiente]} "
        f"{anio_siguiente}: {utm_siguiente:,}".replace(",", ".")
    )
    log_info("[UTM] Excel guardado correctamente.")


# ============================================================
# CORREO - PROYECCIÓN FORMULARIO 29
# ============================================================
def valor_correo(ws, columna: str, fila: int, descripcion: str):
    """Obtiene un valor requerido para el correo sin convertir fórmulas en texto."""
    valor = ws[f"{columna}{fila}"].value
    if isinstance(valor, str) and valor.startswith("="):
        raise AutomatizacionError(
            f"La celda {columna}{fila} ({descripcion}) contiene una fórmula sin calcular."
        )
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        raise AutomatizacionError(
            f"La celda {columna}{fila} ({descripcion}) está vacía."
        )
    return valor


def formatear_monto_correo(valor: object) -> str:
    """Formatea un monto entero con separador de miles chileno, sin anteponer $."""
    if isinstance(valor, bool):
        raise AutomatizacionError(f"Valor monetario inválido para correo: {valor!r}")

    if isinstance(valor, (int, float, Decimal)):
        numero = int(round(float(valor)))
    else:
        texto = str(valor).strip()
        if texto.startswith("="):
            raise AutomatizacionError("Se intentó enviar una fórmula de Excel como monto.")
        numero = monto_entero_chile(texto)

    signo = "-" if numero < 0 else ""
    return signo + f"{abs(numero):,}".replace(",", ".")


def construir_correo_proyeccion(
    ws_valores,
    fila: int,
    anio_actual: int,
    mes_actual: int,
) -> EmailMessage:
    destinatario = texto_excel(
        valor_correo(ws_valores, COL_CORREO, fila, "correo del cliente")
    )
    if "@" not in destinatario or " " in destinatario:
        raise AutomatizacionError(
            f"Correo inválido en {COL_CORREO}{fila}: {destinatario!r}"
        )

    nombre_cliente = texto_excel(
        valor_correo(ws_valores, COL_NOMBRE_CLIENTE, fila, "nombre del cliente")
    )
    empresa = texto_excel(
        valor_correo(ws_valores, COL_RAZON_SOCIAL, fila, "nombre de la empresa")
    )

    iva_debito = formatear_monto_correo(
        valor_correo(ws_valores, COL_IVA_DEBITO, fila, "IVA débito fiscal")
    )
    iva_credito_total = formatear_monto_correo(
        valor_correo(ws_valores, COL_IVA_CREDITO_TOTAL, fila, "IVA crédito fiscal total")
    )
    resultado_iva = texto_excel(
        valor_correo(ws_valores, COL_RESULTADO_IVA, fila, "resultado del IVA")
    )
    iva_por_pagar = formatear_monto_correo(
        valor_correo(ws_valores, COL_IVA_POR_PAGAR, fila, "IVA determinado por pagar")
    )
    remanente_siguiente = formatear_monto_correo(
        valor_correo(
            ws_valores,
            COL_REMANENTE_SIGUIENTE,
            fila,
            "remanente para el período siguiente",
        )
    )
    rbh = formatear_monto_correo(
        valor_correo(ws_valores, COL_RBH, fila, "retención de boletas de honorarios")
    )
    iu = formatear_monto_correo(
        valor_correo(ws_valores, COL_IU, fila, "Impuesto Único de Segunda Categoría")
    )
    ppm = formatear_monto_correo(
        valor_correo(ws_valores, COL_PPM, fila, "PPM")
    )
    total_f29 = formatear_monto_correo(
        valor_correo(ws_valores, COL_TOTAL_F29, fila, "total preliminar del Formulario 29")
    )

    mes_texto = MESES_ES[mes_actual]
    asunto = f"Proyección Formulario 29 {mes_texto} {anio_actual} - {empresa}"

    cuerpo = (
        f"Estimado/a {nombre_cliente}:\n\n"
        f"Junto con saludar, informamos que hemos preparado la proyección preliminar "
        f"del Formulario 29 correspondiente a {mes_texto} de {anio_actual} de {empresa}.\n\n"
        "El resultado estimado es el siguiente:\n\n"
        f"- IVA débito fiscal: ${iva_debito}\n"
        f"- IVA crédito fiscal total: ${iva_credito_total}\n"
        f"- Resultado del IVA: {resultado_iva}\n"
        f"- IVA determinado por pagar: ${iva_por_pagar}\n"
        f"- Remanente para el período siguiente: ${remanente_siguiente}\n"
        f"- Retención de boletas de honorarios: ${rbh}\n"
        f"- Impuesto Único de Segunda Categoría: ${iu}\n"
        f"- PPM: ${ppm}\n\n"
        f"Por lo tanto, el total preliminar del Formulario 29 por pagar asciende a: ${total_f29}\n\n"
        "Agradeceremos revisar esta información y confirmar la autorización para presentar y pagar el formulario.\n\n"
        "Esta comunicación corresponde solamente a una proyección. Los valores podrían modificarse si se "
        "incorporan nuevos documentos o antecedentes antes del cierre definitivo del período.\n\n"
        "Saludos cordiales,"
    )

    html = f"""\
<html>
  <body>
    <p>Estimado/a {nombre_cliente}:</p>
    <p>Junto con saludar, informamos que hemos preparado la proyección preliminar del Formulario 29 correspondiente a <strong>{mes_texto} de {anio_actual}</strong> de <strong>{empresa}</strong>.</p>
    <p>El resultado estimado es el siguiente:</p>
    <ul>
      <li>IVA débito fiscal: <strong>${iva_debito}</strong></li>
      <li>IVA crédito fiscal total: <strong>${iva_credito_total}</strong></li>
      <li>Resultado del IVA: <strong>{resultado_iva}</strong></li>
      <li>IVA determinado por pagar: <strong>${iva_por_pagar}</strong></li>
      <li>Remanente para el período siguiente: <strong>${remanente_siguiente}</strong></li>
      <li>Retención de boletas de honorarios: <strong>${rbh}</strong></li>
      <li>Impuesto Único de Segunda Categoría: <strong>${iu}</strong></li>
      <li>PPM: <strong>${ppm}</strong></li>
    </ul>
    <p>Por lo tanto, el total preliminar del Formulario 29 por pagar asciende a: <strong>${total_f29}</strong></p>
    <p>Agradeceremos revisar esta información y confirmar la autorización para presentar y pagar el formulario.</p>
    <p>Esta comunicación corresponde solamente a una proyección. Los valores podrían modificarse si se incorporan nuevos documentos o antecedentes antes del cierre definitivo del período.</p>
    <p>Saludos cordiales,</p>
  </body>
</html>
"""

    mensaje = EmailMessage()
    mensaje["From"] = CORREO_REMITENTE
    mensaje["To"] = destinatario
    mensaje["Subject"] = asunto
    mensaje.set_content(cuerpo)
    mensaje.add_alternative(html, subtype="html")
    return mensaje


def obtener_servicio_gmail(carpeta_script: Path):
    """
    Crea un servicio autenticado de Gmail API mediante OAuth 2.0.

    - credentials.json: se descarga desde Google Cloud y debe estar junto al script.
    - token.json: se genera automáticamente tras la primera autorización y se reutiliza.
    """
    ruta_credenciales = carpeta_script / ARCHIVO_CREDENCIALES_GOOGLE
    ruta_token = carpeta_script / ARCHIVO_TOKEN_GOOGLE

    if not ruta_credenciales.exists():
        raise AutomatizacionError(
            f"No se encontró {ARCHIVO_CREDENCIALES_GOOGLE} junto al script: "
            f"{ruta_credenciales}"
        )

    credenciales = None

    if ruta_token.exists():
        try:
            credenciales = Credentials.from_authorized_user_file(
                str(ruta_token), GMAIL_SCOPES
            )
        except Exception as exc:
            raise AutomatizacionError(
                f"No se pudo leer {ARCHIVO_TOKEN_GOOGLE}: {exc}"
            ) from exc

    if credenciales and credenciales.expired and credenciales.refresh_token:
        try:
            credenciales.refresh(Request())
        except Exception as exc:
            raise AutomatizacionError(
                "No fue posible renovar la autorización de Gmail. "
                f"Si el token fue revocado o venció, elimina {ARCHIVO_TOKEN_GOOGLE} "
                f"y vuelve a ejecutar para autorizar nuevamente. Detalle: {exc}"
            ) from exc

    if not credenciales or not credenciales.valid:
        try:
            flujo = InstalledAppFlow.from_client_secrets_file(
                str(ruta_credenciales), GMAIL_SCOPES
            )
            credenciales = flujo.run_local_server(port=0)
        except Exception as exc:
            raise AutomatizacionError(
                f"No fue posible completar la autorización OAuth de Gmail: {exc}"
            ) from exc

    try:
        ruta_token.write_text(credenciales.to_json(), encoding="utf-8")
    except Exception as exc:
        raise AutomatizacionError(
            f"No se pudo guardar {ARCHIVO_TOKEN_GOOGLE}: {exc}"
        ) from exc

    try:
        return build("gmail", "v1", credentials=credenciales, cache_discovery=False)
    except Exception as exc:
        raise AutomatizacionError(
            f"No fue posible crear el servicio de Gmail API: {exc}"
        ) from exc


def enviar_mensaje_gmail_api(servicio_gmail, mensaje: EmailMessage) -> dict:
    """Envía un EmailMessage ya construido usando Gmail API."""
    mensaje_raw = base64.urlsafe_b64encode(mensaje.as_bytes()).decode("ascii")
    return (
        servicio_gmail.users()
        .messages()
        .send(userId="me", body={"raw": mensaje_raw})
        .execute()
    )

def filas_con_formulas_correo(ws, filas: Iterable[int]) -> list[str]:
    columnas = [
        COL_IVA_DEBITO,
        COL_IVA_CREDITO_TOTAL,
        COL_RESULTADO_IVA,
        COL_IVA_POR_PAGAR,
        COL_REMANENTE_SIGUIENTE,
        COL_RBH,
        COL_IU,
        COL_PPM,
        COL_TOTAL_F29,
    ]
    formulas: list[str] = []
    for fila in filas:
        for columna in columnas:
            valor = ws[f"{columna}{fila}"].value
            if isinstance(valor, str) and valor.startswith("="):
                formulas.append(f"{columna}{fila}")
    return formulas


def recalcular_excel_con_excel(ruta_excel: Path) -> None:
    """
    Recalcula y guarda el libro con Microsoft Excel mediante COM.
    Solo se usa si las celdas requeridas para los correos contienen fórmulas.
    """
    if sys.platform != "win32":
        raise AutomatizacionError(
            "Las celdas usadas en los correos contienen fórmulas y openpyxl no las calcula. "
            "Se requiere ejecutar este paso en Windows con Microsoft Excel."
        )

    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise AutomatizacionError(
            "Las celdas usadas en los correos contienen fórmulas. Para recalcularlas "
            "automáticamente instala pywin32 con: pip install pywin32"
        ) from exc

    excel = None
    libro = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        libro = excel.Workbooks.Open(str(ruta_excel.resolve()))
        excel.CalculateFullRebuild()
        libro.Save()
    except Exception as exc:
        raise AutomatizacionError(
            f"No fue posible recalcular el Excel antes de enviar los correos: {exc}"
        ) from exc
    finally:
        if libro is not None:
            try:
                libro.Close(SaveChanges=True)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def enviar_correos_proyeccion(
    ruta_excel: Path,
    nombre_hoja: str,
    filas_exitosas: Iterable[int],
    anio_actual: int,
    mes_actual: int,
    carpeta_script: Path,
) -> None:
    filas = list(filas_exitosas)
    if not filas:
        log_info("[CORREO] No hay clientes procesados exitosamente. No se enviarán correos.")
        return

    # Las columnas O/P/Q/R/W/X pueden ser fórmulas de la plantilla. Si lo son,
    # primero forzamos el cálculo real en Microsoft Excel y luego leemos sus resultados.
    wb_formulas = load_workbook(ruta_excel, data_only=False)
    ws_formulas = wb_formulas[nombre_hoja]
    formulas = filas_con_formulas_correo(ws_formulas, filas)
    wb_formulas.close()

    if formulas:
        log_info(
            f"[CORREO] Se detectaron fórmulas en {len(formulas)} celdas requeridas. "
            "Recalculando el libro con Microsoft Excel..."
        )
        recalcular_excel_con_excel(ruta_excel)

    wb_valores = load_workbook(ruta_excel, data_only=True)
    ws_valores = wb_valores[nombre_hoja]

    mensajes: list[tuple[int, EmailMessage]] = []
    for fila in filas:
        try:
            mensaje = construir_correo_proyeccion(
                ws_valores=ws_valores,
                fila=fila,
                anio_actual=anio_actual,
                mes_actual=mes_actual,
            )
            mensajes.append((fila, mensaje))
        except Exception as exc:
            log_error(
                f"[ERROR CORREO] Fila {fila}: no se preparó el correo. "
                f"{type(exc).__name__}: {exc}"
            )

    wb_valores.close()

    if not mensajes:
        log_info("[CORREO] No hay correos válidos para enviar.")
        return

    log_info("[CORREO] Iniciando autenticación OAuth con Gmail API...")
    servicio_gmail = obtener_servicio_gmail(carpeta_script)
    log_info("[CORREO] Gmail API autenticada correctamente.")

    enviados = 0
    for fila, mensaje in mensajes:
        try:
            respuesta = enviar_mensaje_gmail_api(servicio_gmail, mensaje)
            enviados += 1
            id_mensaje = respuesta.get("id", "sin-id")
            log_info(
                f"[CORREO] Fila {fila} enviada correctamente a {mensaje['To']} | "
                f"Asunto: {mensaje['Subject']} | Gmail ID: {id_mensaje}"
            )
        except Exception as exc:
            log_error(
                f"[ERROR CORREO] Fila {fila}: {type(exc).__name__}: {exc}"
            )

    log_info(f"[CORREO] Envío finalizado. Correos enviados: {enviados}/{len(mensajes)}")


# ============================================================
# PROCESAMIENTO DE UN CLIENTE
# ============================================================
def procesar_cliente(
    wb,
    ws,
    ruta_excel: Path,
    fila: int,
    anio_actual: int,
    mes_actual: int,
    carpeta_script: Path,
) -> bool:
    rut_raw = ws[f"{COL_RUT}{fila}"].value
    clave_raw = ws[f"{COL_CLAVE}{fila}"].value

    if not texto_excel(rut_raw):
        log_info(f"[FILA {fila}] Sin RUT. Se omite.")
        return False

    rut = texto_excel(rut_raw)
    nombre = texto_excel(ws[f"{COL_NOMBRE_CLIENTE}{fila}"].value)
    if not nombre:
        nombre = texto_excel(ws[f"{COL_RAZON_SOCIAL}{fila}"].value)
    if not nombre:
        nombre = normalizar_rut(rut)

    imprimir_cliente(fila, rut, nombre)

    clave = texto_excel(clave_raw)
    if not clave:
        log_error("[ERROR] El cliente tiene RUT pero no tiene clave. Se continúa con el siguiente.")
        return False

    driver = None
    try:
        # REQUISITO: navegador nuevo para cada cliente.
        driver = crear_driver(carpeta_script)

        iniciar_sesion(driver, rut, clave)
        log_info("[LOGIN] OK")

        # Compras y Ventas = período actual.
        rcv = procesar_rcv(driver, rut, anio_actual, mes_actual)

        # F29 = mes anterior.
        anio_f29, mes_f29 = periodo_anterior(anio_actual, mes_actual)
        periodo_esperado = periodo_yyyymm(anio_f29, mes_f29)

        pdf = descargar_f29(
            driver=driver,
            carpeta=carpeta_script,
            nombre_cliente=nombre,
            rut=rut,
            anio=anio_f29,
            mes=mes_f29,
        )
        f29 = leer_y_validar_f29(pdf, rut, periodo_esperado)

        # Solo escribimos cuando TODO el cliente terminó correctamente.
        escribir_resultados(ws, fila, rcv, f29)
        guardar_excel(wb, ruta_excel)

        log_info("[EXCEL] H/I/J/S/T/U/V guardados correctamente.")
        log_info("[FINALIZADO]")
        return True

    except NoF29Error as exc:
        log_error(f"[ERROR F29] {exc}")
        log_error("[EXCEL] No se escribieron ceros ni resultados parciales para este cliente.")
        return False
    except Exception as exc:
        log_error(f"[ERROR] {type(exc).__name__}: {exc}")
        log_error("[EXCEL] No se escribieron resultados parciales para este cliente.")
        return False
    finally:
        if driver is not None:
            try:
                driver.quit()
            except Exception:
                pass
        log_info("[CHROME] Sesión cerrada completamente.")


# ============================================================
# ENTRADA DEL PROGRAMA
# ============================================================
def pedir_periodo() -> tuple[int, int]:
    while True:
        entrada = input("Período a procesar (MM/YYYY), ej. 08/2026: ").strip()
        m = re.fullmatch(r"(0[1-9]|1[0-2])/(\d{4})", entrada)
        if m:
            return int(m.group(2)), int(m.group(1))
        print("Formato inválido. Usa MM/YYYY, por ejemplo 08/2026.")


def resolver_excel(carpeta_script: Path) -> Path:
    predeterminado = carpeta_script / ARCHIVO_EXCEL_PREDETERMINADO
    if predeterminado.exists():
        return predeterminado.resolve()

    entrada = input("Nombre o ruta del archivo Excel: ").strip().strip('"')
    ruta = Path(entrada)
    if not ruta.is_absolute():
        ruta = carpeta_script / ruta
    return ruta.resolve()


def main() -> int:
    carpeta_script = Path(__file__).resolve().parent
    configurar_log(carpeta_script)

    try:
        ruta_excel = resolver_excel(carpeta_script)
        anio_actual, mes_actual = pedir_periodo()

        wb, ws = abrir_excel(ruta_excel)

        anio_f29, mes_f29 = periodo_anterior(anio_actual, mes_actual)
        log_info(f"Excel: {ruta_excel}")
        log_info(f"Hoja: {ws.title}")
        log_info(f"Período Compras/Ventas: {mes_actual:02d}/{anio_actual}")
        log_info(f"Período F29: {mes_f29:02d}/{anio_f29}")

        filas_exitosas: list[int] = []
        for fila in range(FILA_INICIAL, ws.max_row + 1):
            procesado_ok = procesar_cliente(
                wb=wb,
                ws=ws,
                ruta_excel=ruta_excel,
                fila=fila,
                anio_actual=anio_actual,
                mes_actual=mes_actual,
                carpeta_script=carpeta_script,
            )
            if procesado_ok:
                filas_exitosas.append(fila)

        # Al finalizar todos los clientes, actualizamos las UTM de D2 y G2
        # usando el mismo período MM/YYYY ingresado al comenzar.
        try:
            actualizar_utm_excel(
                wb=wb,
                ws=ws,
                ruta_excel=ruta_excel,
                anio_actual=anio_actual,
                mes_actual=mes_actual,
            )
        except Exception as exc:
            # Los resultados de clientes exitosos ya fueron guardados.
            # Un problema al consultar UTM no borra ni modifica esos resultados.
            log_error(f"[ERROR UTM] {type(exc).__name__}: {exc}")
            log_error("[UTM] No se actualizaron D2/G2.")

        # Después de completar y guardar la proyección, enviamos los correos solamente
        # para las filas que terminaron correctamente en esta ejecución.
        try:
            enviar_correos_proyeccion(
                ruta_excel=ruta_excel,
                nombre_hoja=NOMBRE_HOJA,
                filas_exitosas=filas_exitosas,
                anio_actual=anio_actual,
                mes_actual=mes_actual,
                carpeta_script=carpeta_script,
            )
        except Exception as exc:
            # Un error de correo no altera los resultados ya guardados en Excel.
            log_error(f"[ERROR CORREO] {type(exc).__name__}: {exc}")
            log_error("[CORREO] Los resultados del Excel se mantienen guardados.")

        log_info("\nProceso terminado.")
        return 0

    except KeyboardInterrupt:
        log_error("\nProceso interrumpido por el usuario.")
        return 130
    except Exception as exc:
        log_error(f"\n[ERROR GENERAL] {type(exc).__name__}: {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
